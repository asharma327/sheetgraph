"""Build tests/fixtures/features.xlsx: one small workbook that uses every xlsx
feature sheetgraph reads beyond values and formulas.

openpyxl writes most of it. The external link and the pivot table are added
afterwards as raw package parts, because openpyxl can't create them.

    pip install openpyxl==3.1.5
    python tests/fixtures/make_features.py
"""

import datetime as dt
import re
import zipfile
import xml.etree.ElementTree as ET
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.styles.colors import Color
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = Path(__file__).with_name("features.xlsx")
REGIONS = ["North", "South", "East", "West"]
PRODUCTS = ["Widget", "Gadget", "Gizmo"]


def build() -> bytes:
    wb = Workbook()

    data = wb.active
    data.title = "Data"
    data.append(["Date", "Region", "Product", "Units", "Price", "Revenue"])
    for i in range(20):
        r = i + 2
        data.append([dt.date(2026, 1, 5) + dt.timedelta(days=7 * i), REGIONS[i % 4], PRODUCTS[i % 3], 10 + i * 7, 2.5 + i, f"=D{r}*E{r}"])
    for c in data[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=Color(theme=4, tint=-0.249977111117893))
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = Border(bottom=Side(style="thin"))
    for row in data.iter_rows(min_row=2, max_row=21):
        row[0].number_format = "yyyy-mm-dd"
        row[4].number_format = "$#,##0.00"
        row[5].number_format = "$#,##0.00"
    data["G1"] = "internal id"
    data.column_dimensions["G"].hidden = True
    data.freeze_panes = "A2"
    data["A1"].comment = Comment("Source: ERP export, weekly", "Analyst")

    sales = Table(displayName="Sales", ref="A1:F21")
    sales.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    data.add_table(sales)

    region = DataValidation(type="list", formula1='"North,South,East,West"', allow_blank=True)
    region.add("B2:B1048576")
    units = DataValidation(type="whole", operator="between", formula1="0", formula2="1000", showErrorMessage=True,
                           promptTitle="Units", prompt="Whole units, 0-1000")
    units.add("D2:D21")
    product = DataValidation(type="list", formula1="Lists!$A$1:$A$3")
    product.add("C2:C21")
    for dv in (region, units, product):
        data.add_data_validation(dv)

    red = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    data.conditional_formatting.add("F2:F21", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="9C0006"), fill=red))
    data.conditional_formatting.add("D2:D21", ColorScaleRule(start_type="min", start_color="F8696B", end_type="max", end_color="63BE7B"))
    data.conditional_formatting.add("A2:F21", FormulaRule(formula=["$D2>150"], font=Font(bold=True)))
    data.conditional_formatting.add("E2:E21", DataBarRule(start_type="min", end_type="max", color="638EC6"))

    lists = wb.create_sheet("Lists")
    for i, p in enumerate(PRODUCTS, 1):
        lists.cell(i, 1, p)
    lists["B1"] = 0.08
    lists["B1"].number_format = "0.0%"
    lists.protection.sheet = True
    wb.defined_names.add(DefinedName("TaxRate", attr_text="Lists!$B$1"))

    summary = wb.create_sheet("Summary")
    summary.sheet_properties.tabColor = "FF0000"
    rows = [
        ("Total revenue", "=SUM(Sales[Revenue])"),
        ("North revenue", '=SUMIF(Data!B:B,"North",Data!F:F)'),
        ("Tax", "=B1*TaxRate"),
        ("FX rate (linked)", "=[1]Rates!B2*2"),
        ("Budget (linked)", "='[1]FX rates'!$C$3+B1"),
    ]
    for r, (label, f) in enumerate(rows, 1):
        summary.cell(r, 1, label)
        summary.cell(r, 2, f)
    summary["A8"] = "hidden row"
    summary.row_dimensions[8].hidden = True
    summary["A10"] = "Docs"
    summary["A10"].hyperlink = "https://example.com/docs"
    summary["A11"] = "Back to data"
    summary["A11"].hyperlink = Hyperlink(ref="A11", location="Data!A1")
    chart = BarChart()
    chart.type = "col"
    chart.title = "Revenue by week"
    chart.add_data(Reference(data, min_col=6, min_row=1, max_row=21), titles_from_data=True)
    chart.set_categories(Reference(data, min_col=1, min_row=2, max_row=21))
    summary.add_chart(chart, "D2")

    log = wb.create_sheet("Log")
    log.append(["Status", "Owner"])
    for row in [("open", "Ana"), ("done", "Ben"), ("open", "Cy"), ("blocked", "Ana")]:
        log.append(row)
    log.auto_filter.ref = "A1:B5"
    log.auto_filter.add_filter_column(0, ["open", "blocked"])

    pivot = wb.create_sheet("Pivot")
    pivot["A1"] = "Product"
    pivot["B1"] = "(All)"
    pivot["A3"] = "Row Labels"
    pivot["B3"] = "Sum of Revenue"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "application/vnd.openxmlformats-officedocument.spreadsheetml"

EXTERNAL_LINK = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<externalLink xmlns="{MAIN}" xmlns:r="{R}"><externalBook r:id="rId1"><sheetNames><sheetName val="Rates"/><sheetName val="FX rates"/></sheetNames><sheetDataSet><sheetData sheetId="0"><row r="2"><cell r="B2"><v>1.1</v></cell></row></sheetData><sheetData sheetId="1"><row r="3"><cell r="C3"><v>1000</v></cell></row></sheetData></sheetDataSet></externalBook></externalLink>"""

EXTERNAL_LINK_RELS = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="{PKG}"><Relationship Id="rId1" Type="{R}/externalLinkPath" Target="file:///C:/Finance/Shared%20Models/Rates.xlsx" TargetMode="External"/></Relationships>"""


def shared(items):
    return f'<sharedItems count="{len(items)}">' + "".join(f'<s v="{v}"/>' for v in items) + "</sharedItems>"


PIVOT_CACHE = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<pivotCacheDefinition xmlns="{MAIN}" xmlns:r="{R}" saveData="0" refreshOnLoad="1" createdVersion="6" refreshedVersion="6" minRefreshableVersion="3" recordCount="0"><cacheSource type="worksheet"><worksheetSource ref="A1:F21" sheet="Data"/></cacheSource><cacheFields count="6"><cacheField name="Date" numFmtId="14"><sharedItems containsSemiMixedTypes="0" containsNonDate="0" containsDate="1" containsString="0"/></cacheField><cacheField name="Region" numFmtId="0">{shared(REGIONS)}</cacheField><cacheField name="Product" numFmtId="0">{shared(PRODUCTS)}</cacheField><cacheField name="Units" numFmtId="0"><sharedItems containsSemiMixedTypes="0" containsString="0" containsNumber="1"/></cacheField><cacheField name="Price" numFmtId="0"><sharedItems containsSemiMixedTypes="0" containsString="0" containsNumber="1"/></cacheField><cacheField name="Revenue" numFmtId="0"><sharedItems containsSemiMixedTypes="0" containsString="0" containsNumber="1"/></cacheField></cacheFields></pivotCacheDefinition>"""

PIVOT_TABLE = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<pivotTableDefinition xmlns="{MAIN}" name="PivotTable1" cacheId="1" dataCaption="Values" updatedVersion="6" minRefreshableVersion="3" createdVersion="6" outline="1" outlineData="1"><location ref="A3:B8" firstHeaderRow="1" firstDataRow="1" firstDataCol="1" rowPageCount="1" colPageCount="1"/><pivotFields count="6"><pivotField showAll="0"/><pivotField axis="axisRow" showAll="0"><items count="5"><item x="0"/><item x="1"/><item x="2"/><item x="3"/><item t="default"/></items></pivotField><pivotField axis="axisPage" showAll="0"><items count="4"><item x="0"/><item x="1"/><item x="2"/><item t="default"/></items></pivotField><pivotField showAll="0"/><pivotField showAll="0"/><pivotField dataField="1" showAll="0"/></pivotFields><rowFields count="1"><field x="1"/></rowFields><pageFields count="1"><pageField fld="2" hier="-1"/></pageFields><dataFields count="1"><dataField name="Sum of Revenue" fld="5" baseField="0" baseItem="0"/></dataFields></pivotTableDefinition>"""

PIVOT_TABLE_RELS = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="{PKG}"><Relationship Id="rId1" Type="{R}/pivotCacheDefinition" Target="../pivotCache/pivotCacheDefinition1.xml"/></Relationships>"""


def rel(rid, kind, target):
    return f'<Relationship Id="{rid}" Type="{R}/{kind}" Target="{target}"/>'


def add_parts(xlsx: bytes) -> bytes:
    with zipfile.ZipFile(BytesIO(xlsx)) as z:
        files = {n: z.read(n) for n in z.namelist()}

    workbook = files["xl/workbook.xml"].decode()
    wb_rels = files["xl/_rels/workbook.xml.rels"].decode()
    types = files["[Content_Types].xml"].decode()

    # Which worksheet part is "Pivot"?
    rid = re.search(r'<sheet\b[^>]*\bname="Pivot"[^>]*\br:id="([^"]+)"', workbook).group(1)
    target = next(r.get("Target") for r in ET.fromstring(wb_rels) if r.get("Id") == rid)
    sheet_part = target.lstrip("/") if target.startswith("/") else "xl/" + target
    folder, name = sheet_part.rsplit("/", 1)
    sheet_rels = f"{folder}/_rels/{name}.rels"

    workbook = workbook.replace("</sheets>", '</sheets><externalReferences><externalReference r:id="rIdExt1"/></externalReferences>', 1)
    workbook = workbook.replace("</workbook>", '<pivotCaches><pivotCache cacheId="1" r:id="rIdPc1"/></pivotCaches></workbook>', 1)
    wb_rels = wb_rels.replace(
        "</Relationships>",
        rel("rIdExt1", "externalLink", "externalLinks/externalLink1.xml")
        + rel("rIdPc1", "pivotCacheDefinition", "pivotCache/pivotCacheDefinition1.xml")
        + "</Relationships>",
    )
    types = types.replace(
        "</Types>",
        f'<Override PartName="/xl/externalLinks/externalLink1.xml" ContentType="{CT}.externalLink+xml"/>'
        f'<Override PartName="/xl/pivotCache/pivotCacheDefinition1.xml" ContentType="{CT}.pivotCacheDefinition+xml"/>'
        f'<Override PartName="/xl/pivotTables/pivotTable1.xml" ContentType="{CT}.pivotTable+xml"/>'
        "</Types>",
    )
    pivot_rel = rel("rIdPt1", "pivotTable", "../pivotTables/pivotTable1.xml")
    if sheet_rels in files:
        files[sheet_rels] = files[sheet_rels].decode().replace("</Relationships>", pivot_rel + "</Relationships>").encode()
    else:
        files[sheet_rels] = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="{PKG}">{pivot_rel}</Relationships>'.encode()

    files.update({
        "xl/workbook.xml": workbook.encode(),
        "xl/_rels/workbook.xml.rels": wb_rels.encode(),
        "[Content_Types].xml": types.encode(),
        "xl/externalLinks/externalLink1.xml": EXTERNAL_LINK.encode(),
        "xl/externalLinks/_rels/externalLink1.xml.rels": EXTERNAL_LINK_RELS.encode(),
        "xl/pivotCache/pivotCacheDefinition1.xml": PIVOT_CACHE.encode(),
        "xl/pivotTables/pivotTable1.xml": PIVOT_TABLE.encode(),
        "xl/pivotTables/_rels/pivotTable1.xml.rels": PIVOT_TABLE_RELS.encode(),
    })
    out = BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        # [Content_Types].xml first, as Excel writes it.
        for n in sorted(files, key=lambda n: (n != "[Content_Types].xml", n)):
            z.writestr(n, files[n])
    return out.getvalue()


if __name__ == "__main__":
    OUT.write_bytes(add_parts(build()))
    print(f"wrote {OUT}")
