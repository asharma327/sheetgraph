"""Messy small-business job tracker for the sheetgraph demo image."""
import random
import sys
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

random.seed(7)
out = sys.argv[1]
N = 12000

wb = Workbook()
jobs = wb.active
jobs.title = "Jobs"
rates = wb.create_sheet("Rates")
summary = wb.create_sheet("Summary")
wb.create_sheet("DO NOT TOUCH")

subs = ["Priya", "Dmitri K", "Sam", "Tom H", "Maya"]
trades = ["deck", "garage", "siding", "roof", "kitchen", "bath"]
rates.append(["Trade", "Rate", "Markup"])
for t in trades:
    rates.append([t, random.choice([85, 95, 110, 125]), random.choice([0.15, 0.2, 0.25])])

jobs.append(["Job #", "Address", "Start", "Sub", "Trade", "Hours", "Rate", "Quote", "Paid?", "Balance", "Month", "Notes"])
for c in jobs[1]:
    c.font = Font(bold=True)
streets = ["Wexford Ln", "Fairmont St", "Larkspur Dr", "Cedar Rise", "Beacon Hill", "Marlow Cl", "Saltmarsh", "Harrow St"]
start = date(2021, 1, 4)
yellow, pink = PatternFill("solid", fgColor="FFFF00"), PatternFill("solid", fgColor="FF99CC")
for i in range(2, N + 2):
    jobs.append([
        f"J-{i:05d}",
        f"{random.randint(1, 400)} {random.choice(streets)}",
        start + timedelta(days=i // 6),
        random.choice(subs),
        random.choice(trades),
        random.randint(4, 120),
        f"=VLOOKUP(E{i},Rates!$A$2:$C$7,2,FALSE)",
        f"=F{i}*G{i}*(1+VLOOKUP(E{i},Rates!$A$2:$C$7,3,FALSE))",
        random.choice(["Y", "N", "partial"]),
        f'=IF(I{i}="Y",0,H{i})',
        f"=MONTH(C{i})",
        random.choice(["", "", "waiting on permit", "quote sent", "no answer", "deck only", "cancelled?? ask maya"]),
    ])
    if random.random() < 0.02:  # someone typed over the formula
        jobs[f"H{i}"] = random.randint(5000, 90000)
    if random.random() < 0.05:
        jobs[f"A{i}"].fill = random.choice([yellow, pink])
dv = DataValidation(type="list", formula1='"Y,N,partial"')
jobs.add_data_validation(dv)
dv.add(f"I2:I{N + 1}")

summary.append(["Month", "Quoted", "Outstanding"])
for m in range(1, 13):
    summary.append([m, f"=SUMIF(Jobs!K:K,A{m + 1},Jobs!H:H)", f"=SUMIF(Jobs!K:K,A{m + 1},Jobs!J:J)"])
summary.append(["Total", "=SUM(B2:B13)", "=SUM(C2:C13)"])
summary["E1"] = "=3820+3770+1250"
wb.save(out)
print(out)
